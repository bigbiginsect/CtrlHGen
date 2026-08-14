#!/usr/bin/env python3
"""Insert three planned-improvement slides into the current 15-slide deck.

This intentionally treats the manually edited PPTX as the source of truth.
Slides 1-14 (including their notes) are copied byte-for-byte into the final
archive after python-pptx creates the three new slide parts.
"""

from __future__ import annotations

import hashlib
import shutil
import sys
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import build_ctrlhgen_teacher_briefing_minimal as style  # noqa: E402


PPTX = ROOT / "worklogs/reproduction/briefings/CtrlHGen-论文讲解与复现进展-2026-08-13-极简重制版.pptx"
ROUNDTRIP = Path("/tmp/ctrlhgen-improvements-roundtrip.pptx")
STAGED = Path("/tmp/ctrlhgen-improvements-final.pptx")


NOTES = [
    {
        "time": "1:00",
        "talk": (
            "前面已经完成论文方法和复现结果的汇报。这里开始讲尚未实施的改进计划。"
            "论文把 Dice 和 Overlap 称为平滑语义奖励，但 Dice 可以写成 2J/(1+J)，"
            "因此与 Jaccard 的候选排序完全相同；Overlap 又会在两个答案集存在包含关系时接近 1。"
            "更关键的是，specific control 只检查条件 token 是否出现，无法判断它是否真正参与解释。"
            "本地 rollout 中有 60.55% 的组奖励方差为零，而只有 37.89% 的组四条文本完全相同，"
            "两者相差 22.66 个百分点，说明确实存在不同逻辑选择却无法被现有奖励区分的部分。"
        ),
        "transition": "第一个备选方向是直接利用逻辑查询可执行这一特点，测量每个槽位的反事实贡献。",
        "caveat": (
            "22.66 个百分点是两个组级比例的差值，用于定位奖励盲区，不代表加入新奖励后一定能获得同等幅度的训练收益。"
            "以下三页全部是计划设计，不是已经得到的实验结果。"
        ),
        "qa": (
            "问：Dice 至少是否让奖励更平滑？答：它会重缩放数值间距，但不改变候选排序，"
            "而且仍是完整逻辑执行后的终局标量，因此没有提供槽位级信用分配。"
        ),
    },
    {
        "time": "1:15",
        "talk": (
            "方案一对生成假设中的关系或实体槽位做匹配替换：保持 pattern、槽位数量和变量绑定不变，"
            "并尽量匹配关系方向、类型与度数。原假设得分减去替换后的期望得分，就是该槽位的反事实边际贡献。"
            "对 specific condition，再把条件是否出现、完整假设能否解释观察，以及替换条件后答案是否改变同时纳入 effective adherence。"
            "这样，BAFTA 式无效 OR 分支虽然名义遵循为 1，但因为替换或移除后结论不变，实际贡献会接近 0。"
            "首轮不训练，只在已有 256×4 rollouts 上离线重算，并构造 OR-append 对抗对。"
        ),
        "transition": "另一条独立路线不改奖励，而是利用 pattern 已经给出的完整逻辑计划，直接压缩生成空间。",
        "caveat": (
            "结果高度依赖 matched replacement 分布；同义关系可能具有真实贡献却在替换后变化很小。"
            "每个 rollout 还会增加约 2 到 K 次图执行，完全相同的 completion 也仍然无法靠 IDC 打破平局。"
        ),
        "qa": (
            "问：为什么不用直接删除槽位？答：删除往往同时改变逻辑复杂度和变量结构；"
            "匹配替换更接近只改变指称选择的受控反事实。"
        ),
    },
    {
        "time": "1:15",
        "talk": (
            "方案二把 pattern condition 编译成固定模板，其中 i、u、n 和 END 等操作符由编译器决定，"
            "REL 位置只允许关系 token，ENT 位置只允许实体 token。模型不再重新生成输入中已经给出的 logical plan，"
            "而只学习 observation 和 condition 下的 KG grounding。最小实验仍不重新训练：在同一 checkpoint、"
            "同一验证集和相同采样种子下比较 free decoding、通用 syntax FSA 与 condition-template FSA。"
            "主要看 Jaccard、Dice、Overlap 和推理开销；parse rate 与 Pattern Accuracy 因为被机制保证，不能再算作模型能力提升。"
        ),
        "transition": "两个方案先平行、独立证伪；只有各自通过后，才考虑 Decoder 与 IDC 的组合实验。",
        "caveat": (
            "首轮 template FSA 只针对当前 pattern control；不能据此外推 specific entity、specific relation 或其他控制类型。"
            "受约束后 PA 和 parse 的提升属于搜索机制效果，checkpoint 选择必须转向集合语义指标。"
        ),
        "qa": (
            "问：为什么不直接把两项改动一起训练？答：先独立实验才能判断收益来自搜索空间收缩还是反事实奖励；"
            "两条路线都成立后，再做组合才有清晰的因果解释。"
        ),
    },
]

SOURCES = [
    "akgr/abduction_model/reproduction.py:335–344；akgr/evaluation.py:314–320；worklogs/reproduction/report-data/phase-d/repaired-full/rollout-signal-audit.json；论文 Fig. 9",
    "akgr/tokenizer.py:56–76；akgr/evaluation.py:314–320；论文 Fig. 9；现有 256×4 rollouts",
    "akgr/abduction_model/experiment_runner.py:167–185；pattern condition action format；现有 SFT/GRPO checkpoint",
]


def add_shell(prs: Presentation, number: int, title: str, section: str, note: dict, source: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = style.WHITE
    style.add_header(slide, number, title, section)
    style.add_line(slide, 0.68, 6.88, 12.65, 6.88, color=style.LIGHT, width=0.6)
    style.add_notes(slide, note, source)
    return slide


def add_metric(slide, x: float, value: str, label: str, color=style.BLUE) -> None:
    style.add_text(slide, x, 5.03, 2.74, 0.34, value, size=21, color=color, bold=True, align=PP_ALIGN.CENTER)
    style.add_text(slide, x, 5.48, 2.74, 0.42, label, size=11.8, color=style.DARK, align=PP_ALIGN.CENTER)


def add_flow_box(slide, x: float, w: float, title: str, detail: str, emphasis=False) -> None:
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(x), Inches(1.35), Inches(w), Inches(1.18),
    )
    shape.adjustments[0] = 0.06
    shape.fill.solid()
    shape.fill.fore_color.rgb = style.LIGHT_BLUE if emphasis else style.VERY_LIGHT
    shape.line.color.rgb = style.BLUE if emphasis else style.LIGHT
    shape.line.width = Pt(1.0)
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.08)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = title
    style.set_run_font(r, 15.5, color=style.BLUE if emphasis else style.BLACK, bold=True)
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER; p2.space_before = Pt(5)
    r2 = p2.add_run(); r2.text = detail
    style.set_run_font(r2, 11.5, color=style.DARK)


def add_diagnosis_slide(prs: Presentation) -> None:
    slide = add_shell(
        prs, 15, "为什么还需要改进：终局奖励没有解决信用分配",
        "Next step · Diagnosis", NOTES[0], SOURCES[0],
    )
    style.add_text(slide, 0.82, 1.28, 5.35, 0.34, "1  排序没有增加信息", size=17, bold=True)
    style.add_text(slide, 1.05, 1.84, 4.85, 0.42, "Dice = 2J / (1 + J)", size=25, color=style.BLUE, bold=True, align=PP_ALIGN.CENTER)
    style.add_text(slide, 1.02, 2.43, 4.92, 0.54,
                   "对 J 严格单调递增\n候选排序与 Jaccard 完全相同，只是重缩放间距",
                   size=14.1, color=style.DARK, align=PP_ALIGN.CENTER)
    style.add_text(slide, 0.96, 3.28, 5.02, 0.38,
                   "Overlap = |A ∩ O| / min(|A|, |O|)", size=18, color=style.BLACK, bold=True, align=PP_ALIGN.CENTER)
    style.add_text(slide, 1.02, 3.82, 4.92, 0.52,
                   "A ⊇ O 或 O ⊇ A 时可取 1\n→ 严重过覆盖 / 欠覆盖仍可能被奖励",
                   size=13.8, color=style.RED, align=PP_ALIGN.CENTER)

    style.add_line(slide, 6.40, 1.28, 6.40, 4.52, color=style.LIGHT, width=0.8)
    style.add_text(slide, 6.78, 1.28, 5.45, 0.34, "2  “出现”不等于“参与解释”", size=17, bold=True)
    style.add_text(slide, 7.05, 1.93, 4.95, 0.38, "H = H_effective   ∨   H_BAFTA", size=22, color=style.BLUE, bold=True, align=PP_ALIGN.CENTER)
    style.add_text(slide, 7.22, 2.63, 4.62, 0.34, "[[H]]_G = [[H_effective]]_G = O", size=17.5, color=style.BLACK, bold=True, align=PP_ALIGN.CENTER)
    style.add_text(slide, 7.22, 3.18, 4.62, 0.44,
                   "替换或删除 BAFTA 分支\n结论集合完全不变", size=14.2, color=style.DARK, align=PP_ALIGN.CENTER)
    style.add_rich(slide, 6.88, 3.98, 5.28, 0.40, [
        ("nominal adherence = 1", style.RED, True, False),
        ("     |     ", style.GRAY, False, False),
        ("effective contribution ≈ 0", style.BLUE, True, False),
    ], size=12.2, align=PP_ALIGN.CENTER)

    style.add_line(slide, 0.82, 4.76, 12.22, 4.76, color=style.LIGHT, width=0.8)
    for x in (3.55, 6.44, 9.33):
        style.add_line(slide, x, 4.96, x, 5.91, color=style.LIGHT, width=0.6)
    add_metric(slide, 0.82, "95.43%", "Pattern Accuracy")
    add_metric(slide, 3.71, "98.24%", "condition accuracy")
    add_metric(slide, 6.60, "60.55%", "zero-reward-variance groups", color=style.RED)
    add_metric(slide, 9.49, "37.89%", "all strings identical")
    style.add_line(slide, 0.84, 6.19, 0.84, 6.55, color=style.BLUE, width=2.4)
    style.add_text(slide, 1.02, 6.17, 11.25, 0.39,
                   "60.55 − 37.89 = 22.66 pp：文本不同，现有终局奖励仍然相同。",
                   size=15.0, color=style.BLUE, bold=True)


def add_idc_slide(prs: Presentation) -> None:
    slide = add_shell(
        prs, 16, "备选方案一：槽位对称的反事实贡献奖励",
        "Next step · Option 1", NOTES[1], SOURCES[1],
    )
    style.add_text(slide, 0.82, 1.25, 7.35, 0.32, "Counterfactual marginal contribution", size=16.5, bold=True)
    style.add_text(slide, 0.92, 1.78, 7.18, 0.46,
                   "Δ_z = S([[H]]_G, O) − E_{z′ ~ q_match} S([[H_{z←z′}]]_G, O)",
                   size=18.5, color=style.BLUE, bold=True, align=PP_ALIGN.CENTER)
    style.add_text(slide, 0.92, 2.45, 7.18, 0.86,
                   "g_eff(H,C) = 1[C∈H] · S([[H]]_G,O)\n"
                   "× E_{C′}[1 − J([[H]]_G, [[H_{C←C′}]]_G)]",
                   size=17.0, color=style.BLACK, bold=True, align=PP_ALIGN.CENTER)
    style.add_rich(slide, 0.98, 3.48, 7.08, 0.38, [
        ("BAFTA 式伪遵循：", style.BLACK, True, False),
        ("条件出现", style.RED, True, False),
        ("，但替换后答案不变  ⇒  ", style.BLACK, False, False),
        ("g_eff ≈ 0", style.BLUE, True, False),
    ], size=13.6, align=PP_ALIGN.CENTER)

    style.add_line(slide, 8.35, 1.30, 8.35, 3.93, color=style.LIGHT, width=0.8)
    style.add_text(slide, 8.72, 1.28, 3.56, 0.34, "matched replacement", size=16.5, bold=True)
    style.add_bullets(slide, 8.72, 1.86, 3.52, 1.46, [
        "pattern、槽位数、变量绑定不变",
        "relation 方向、domain/range、度数尽量匹配",
        "entity 类型、语义角色、邻域度数尽量匹配",
    ], size=12.8, gap=7)
    style.add_text(slide, 8.72, 3.42, 3.42, 0.50,
                   "前置修正：所有 unique eligible slots\n均匀采样（baseline hygiene）",
                   size=11.8, color=style.RED, bold=True, align=PP_ALIGN.CENTER)

    style.add_line(slide, 0.82, 4.15, 12.22, 4.15, color=style.LIGHT, width=0.8)
    style.add_text(slide, 0.84, 4.39, 11.5, 0.31, "最小证伪：先离线，不训练", size=16.5, color=style.BLUE, bold=True)
    steps = [
        (0.90, "1", "已有 256×4 rollouts", "离线重算 IDC"),
        (4.42, "2", "原奖励  vs  +IDC", "统计非同字符串等奖率"),
        (8.18, "3", "BAFTA OR-append 对抗对", "比较真实贡献与伪分支"),
    ]
    for x, number, title, detail in steps:
        style.add_text(slide, x, 4.94, 0.36, 0.34, number, size=17, color=style.BLUE, bold=True, align=PP_ALIGN.CENTER)
        style.add_text(slide, x + 0.48, 4.92, 3.15, 0.31, title, size=13.2, bold=True)
        style.add_text(slide, x + 0.48, 5.35, 3.15, 0.30, detail, size=12.2, color=style.DARK)
    style.add_text(slide, 0.96, 5.86, 11.35, 0.26,
                   "成本：每个 rollout 额外约 2–K 次图执行；主要风险是匹配分布与同义关系。",
                   size=11.2, color=style.GRAY, align=PP_ALIGN.CENTER)
    style.add_line(slide, 0.84, 6.23, 0.84, 6.58, color=style.BLUE, width=2.4)
    style.add_text(slide, 1.02, 6.20, 11.25, 0.40,
                   "Go：等奖率的配对 95% CI < 0，且真实条件贡献显著高于伪分支；否则停止。",
                   size=13.8, color=style.BLUE, bold=True)


def add_decoder_slide(prs: Presentation) -> None:
    slide = add_shell(
        prs, 17, "备选方案二：Pattern-Compiled Grounding Decoder",
        "Next step · Option 2", NOTES[2], SOURCES[2],
    )
    add_flow_box(slide, 0.78, 2.00, "pattern C", "已知 logical plan")
    style.add_text(slide, 2.83, 1.70, 0.42, 0.35, "→", size=24, color=style.BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_flow_box(slide, 3.26, 3.82, "τ(C) = [ i, REL, ENT,\nn, REL, ENT, END ]", "i / u / n / END 由编译器固定", emphasis=True)
    style.add_text(slide, 7.12, 1.70, 0.42, 0.35, "→", size=24, color=style.BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_flow_box(slide, 7.55, 2.18, "typed masks", "REL 只选关系 · ENT 只选实体")
    style.add_text(slide, 9.78, 1.70, 0.42, 0.35, "→", size=24, color=style.BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_flow_box(slide, 10.20, 2.30, "Instantiate H", "模型只搜索 grounding")

    style.add_text(slide, 0.82, 2.86, 11.45, 0.32, "同一 checkpoint 的零训练三路对比", size=16.5, color=style.BLUE, bold=True)
    rows = [
        ("解码", "逻辑结构", "token 可行域", "模型实际搜索"),
        ("Free decoding", "自由生成", "全词表", "完整操作符 + grounding 序列"),
        ("Syntax FSA", "保证语法合法", "通用 token 类", "合法但仍完整的逻辑序列"),
        ("Template FSA", "由 condition 固定", "REL / ENT typed masks", "仅实体与关系槽位"),
    ]
    table = slide.shapes.add_table(4, 4, Inches(0.82), Inches(3.34), Inches(11.52), Inches(2.02)).table
    widths = [2.00, 2.50, 2.85, 4.17]
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(width)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            fill = style.VERY_LIGHT if i == 0 else (style.LIGHT_BLUE if i == 3 else style.WHITE)
            style.set_cell(
                table.cell(i, j), value, size=12.5 if i else 12.8,
                color=style.BLUE if i == 3 else style.BLACK,
                bold=(i == 0 or i == 3), fill=fill,
            )
    style.add_rich(slide, 0.98, 5.58, 11.30, 0.32, [
        ("主指标：", style.BLACK, True, False),
        ("Jaccard / Dice / Overlap + latency", style.BLUE, True, False),
        ("       机制指标：", style.BLACK, True, False),
        ("Parse / PA（保证，不计为能力提升）", style.RED, True, False),
    ], size=12.6, align=PP_ALIGN.CENTER)
    style.add_line(slide, 0.84, 6.04, 0.84, 6.40, color=style.BLUE, width=2.4)
    style.add_text(slide, 1.02, 6.02, 11.22, 0.38,
                   "Go：Template 对 Free 与 Syntax 的 paired Δsemantic 95% CI 均 > 0 → 再做 slot-only / type-masked SFT。",
                   size=12.9, color=style.BLUE, bold=True)
    style.add_text(slide, 1.02, 6.46, 11.20, 0.20,
                   "两个方案先独立验证；均通过后，再考虑 Decoder × IDC。",
                   size=10.8, color=style.GRAY, align=PP_ALIGN.CENTER)


def title_of(slide) -> str:
    shapes = sorted(
        [shape for shape in slide.shapes if getattr(shape, "has_text_frame", False) and shape.text.strip()],
        key=lambda shape: (shape.top, shape.left),
    )
    return shapes[0].text.splitlines()[0] if shapes else ""


def update_backup_page_number(slide) -> None:
    candidates = [
        shape for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
        and shape.text.strip() == "17"
        and shape.left > Inches(12)
        and shape.top > Inches(6.8)
    ]
    if len(candidates) != 1:
        raise AssertionError(f"expected one backup page-number shape, found {len(candidates)}")
    runs = candidates[0].text_frame.paragraphs[0].runs
    if len(runs) != 1:
        raise AssertionError("unexpected page-number run structure")
    runs[0].text = "18"


def merge_preserving_original_slides(original: Path, roundtrip: Path, staged: Path) -> None:
    mutable_common = {
        "[Content_Types].xml",
        "docProps/app.xml",
        "ppt/presentation.xml",
        "ppt/_rels/presentation.xml.rels",
        "ppt/slides/slide15.xml",
    }
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(roundtrip) as generated:
        source_names = set(source.namelist())
        with zipfile.ZipFile(staged, "w") as output:
            for info in generated.infolist():
                name = info.filename
                if name in source_names and name not in mutable_common:
                    output.writestr(source.getinfo(name), source.read(name))
                else:
                    output.writestr(info, generated.read(name))


def main() -> None:
    before = zipfile.ZipFile(PPTX)
    before_slide_hashes = {
        index: hashlib.sha256(before.read(f"ppt/slides/slide{index}.xml")).hexdigest()
        for index in range(1, 15)
    }
    before_note_hashes = {
        index: hashlib.sha256(before.read(f"ppt/notesSlides/notesSlide{index}.xml")).hexdigest()
        for index in range(1, 15)
    }
    before.close()

    prs = Presentation(PPTX)
    if len(prs.slides) != 15:
        raise AssertionError(f"expected current 15-slide source deck, got {len(prs.slides)}")
    original_titles = [title_of(prs.slides[index]) for index in range(14)]
    original_slide14_notes = prs.slides[13].notes_slide.notes_text_frame.text
    backup_slide = prs.slides[14]
    update_backup_page_number(backup_slide)

    add_diagnosis_slide(prs)
    add_idc_slide(prs)
    add_decoder_slide(prs)

    # New parts are appended as slide16-18. Reorder them before the original
    # slide15 while keeping slide1-14 exactly where they were.
    slide_ids = prs.slides._sldIdLst
    new_ids = list(slide_ids[-3:])
    for slide_id in new_ids:
        slide_ids.remove(slide_id)
    for offset, slide_id in enumerate(new_ids):
        slide_ids.insert(14 + offset, slide_id)

    prs.save(ROUNDTRIP)
    merge_preserving_original_slides(PPTX, ROUNDTRIP, STAGED)

    if zipfile.ZipFile(STAGED).testzip() is not None:
        raise AssertionError("generated PPTX archive failed CRC validation")
    check = Presentation(STAGED)
    if len(check.slides) != 18:
        raise AssertionError(f"expected 18 slides, got {len(check.slides)}")
    expected_titles = original_titles + [
        "为什么还需要改进：终局奖励没有解决信用分配",
        "备选方案一：槽位对称的反事实贡献奖励",
        "备选方案二：Pattern-Compiled Grounding Decoder",
        "备份：完整双解码指标",
    ]
    actual_titles = [title_of(slide) for slide in check.slides]
    if actual_titles != expected_titles:
        raise AssertionError((actual_titles, expected_titles))
    if check.slides[13].notes_slide.notes_text_frame.text != original_slide14_notes:
        raise AssertionError("slide 14 notes changed")
    if not any(
        getattr(shape, "has_text_frame", False) and shape.text.strip() == "18"
        for shape in check.slides[17].shapes
    ):
        raise AssertionError("backup slide page number was not updated to 18")
    for index in range(14, 17):
        slide = check.slides[index]
        notes = slide.notes_slide.notes_text_frame.text
        for marker in ("【建议用时】", "【核心讲述】", "【转场】", "【证据来源】", "【必要限定】", "【可能追问】"):
            if marker not in notes:
                raise AssertionError((title_of(slide), marker))

    with zipfile.ZipFile(STAGED) as final_zip:
        for index in range(1, 15):
            slide_hash = hashlib.sha256(final_zip.read(f"ppt/slides/slide{index}.xml")).hexdigest()
            note_hash = hashlib.sha256(final_zip.read(f"ppt/notesSlides/notesSlide{index}.xml")).hexdigest()
            if slide_hash != before_slide_hashes[index] or note_hash != before_note_hashes[index]:
                raise AssertionError(f"original slide or notes changed at index {index}")

    try:
        shutil.copy2(STAGED, PPTX)
        output = PPTX
    except PermissionError:
        output = PPTX.with_name(f"{PPTX.stem}-含改进方案.pptx")
        shutil.copy2(STAGED, output)
        print(f"Source deck is locked; wrote {output} instead")
    else:
        print(f"Updated {output}")
    print("Slides: 18; original slides 1-14 and their notes preserved byte-for-byte")


if __name__ == "__main__":
    main()
